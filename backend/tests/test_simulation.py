from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
import uuid

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas.equipment import EquipmentInput
from app.schemas.economics import EconomicsInput
from app.schemas.simulation import SimulationOptions, SimulationRequest
from app.services.equipment import equipment_plan
from app.services.economics import calculate
from app.services.economics_reports import workbook_report
from app.services.simulation import simulate


def snapshot(**equipment):
    return {"name": "Тестовый процесс", "quantity": 3,
        "context": {"hours_per_day": 8, "daily_demand": 400, "peak_factor": 1, "unit": "паллет/ч", "missing": []},
        "candidate": {"throughput": 30, "status": "needs_review", "missing": ["Проверить пол"]},
        "selection_options": {"utilization": .8, "availability": .9},
        "equipment": EquipmentInput(peak_rate=50, **equipment).model_dump()}


def test_determinism_conservation_and_complete_robot_timelines():
    source = snapshot(runtime_hours=.5, charge_hours=.2)
    original = deepcopy(source)
    result = simulate(source, SimulationOptions())
    assert result == simulate(source, SimulationOptions())
    assert source == original
    kpi = result["kpi"]
    assert kpi["completed"] + kpi["backlog"] == pytest.approx(kpi["target"])
    assert kpi["state_hours"]["charging"] > 0
    assert sum(kpi["state_hours"].values()) == pytest.approx(24, abs=.001)
    for robot in result["robots"]:
        segments = robot["segments"]
        assert segments[0]["start"] == 0
        assert segments[-1]["end"] == result["duration_seconds"]
        for a, b in zip(segments, segments[1:]):
            assert a["end"] == pytest.approx(b["start"])
    assert result["frames"][-1]["completed"] == kpi["completed"]
    assert all(a["completed"] <= b["completed"] for a, b in zip(result["frames"], result["frames"][1:]))


def test_fewer_workstations_cause_measurable_bottleneck():
    few = snapshot(handling_seconds=180, station_count=1, override_reason="Один доступный пост")
    many = snapshot(handling_seconds=180, station_count=3, override_reason="Три доступных поста")
    a, b = [simulate(s, SimulationOptions(route_m=0)) for s in (few, many)]
    assert a["kpi"]["completed"] < b["kpi"]["completed"]
    assert a["kpi"]["state_hours"]["station_queue"] > b["kpi"]["state_hours"]["station_queue"]


def test_more_chargers_reduce_charging_queue():
    few = snapshot(runtime_hours=.2, charge_hours=.2, charger_count=1, station_count=3, override_reason="Одна зарядка, три поста")
    many = snapshot(runtime_hours=.2, charge_hours=.2, charger_count=3, station_count=3, override_reason="Три зарядки, три поста")
    a, b = [simulate(s, SimulationOptions(route_m=0)) for s in (few, many)]
    assert a["kpi"]["completed"] < b["kpi"]["completed"]
    assert a["kpi"]["state_hours"]["charger_queue"] > b["kpi"]["state_hours"]["charger_queue"]


def test_long_route_invalidates_optimistic_selection_capacity():
    source = snapshot()
    short = simulate(source, SimulationOptions(route_m=0))
    long = simulate(source, SimulationOptions(route_m=400))
    assert short["kpi"]["completed"] > long["kpi"]["completed"]
    assert long["kpi"]["throughput"] < long["kpi"]["selection_capacity"]
    assert any("Маршрут" in w for w in long["warnings"])


def test_fractional_batch_does_not_invent_demand():
    source = snapshot(operations_per_cycle=7)
    source["context"]["daily_demand"] = 10.5
    result = simulate(source, SimulationOptions(route_m=0))
    assert result["kpi"]["target"] == 10.5
    assert result["kpi"]["completed"] <= 10.5
    assert result["frames"][-1]["arrived"] == 10.5


@pytest.mark.parametrize("change", ["fleet", "battery", "excluded"])
def test_unsupported_workload_fails_with_actionable_error(change):
    source = snapshot()
    if change == "fleet": source["quantity"] = 201
    if change == "battery": source["equipment"]["runtime_hours"] = .001
    if change == "excluded": source["candidate"]["status"] = "excluded"
    with pytest.raises(ValueError):
        simulate(source, SimulationOptions())


def test_sizing_overrides_prices_and_recalculation_for_changed_fleet():
    inputs = EquipmentInput(peak_rate=120, runtime_hours=8, charge_hours=2,
        handling_seconds=60, charger_price=100, station_price=300)
    small = equipment_plan(4, inputs)
    large = equipment_plan(20, inputs)
    assert [i["quantity"] for i in small["items"]] == [1, 3]
    assert small["total_cost"] == 1000
    assert large["items"][0]["quantity"] == 5
    manual = equipment_plan(4, inputs.model_copy(update={"charger_count": 2}))
    assert manual["items"][0]["calculated"] == 1
    assert manual["items"][0]["quantity"] == 2
    with pytest.raises(ValidationError): EquipmentInput(charger_count=2)


def test_equipment_costs_flow_to_both_economic_modes_and_export():
    equipment = EquipmentInput(peak_rate=120, handling_seconds=60, charger_price=100, station_price=300)
    data = EconomicsInput(baseline_annual_labor=10000, scenarios=[
        {"name": "Покупка", "mode": "purchase", "quantity": 4, "equipment_price": 1000, "equipment": equipment},
        {"name": "RaaS", "mode": "raas", "quantity": 4, "monthly_fee": 10, "equipment": equipment}])
    result = calculate(data)
    assert result.results[0].capex == 5000
    assert result.results[1].capex == 1000
    assert result.results[1].annual_opex == 10480
    assert result.results[0].tco == 55000
    book = load_workbook(BytesIO(workbook_report(result.model_dump(mode="json"))))
    assert book["Состав оборудования"].max_row == 5
    assert book["Состав оборудования"]["F2"].value == 100


def test_request_rejects_mixed_sources_and_nonfinite_options():
    with pytest.raises(ValidationError):
        SimulationRequest(project_updated_at=datetime.now(UTC), economics_run_id=uuid.uuid4(), product_id=1)
    with pytest.raises(ValidationError): SimulationOptions(route_m=float("nan"))


@pytest.mark.asyncio
async def test_saved_source_is_scoped_to_project_and_uses_immutable_snapshot():
    from app.api.routes.simulation import simulation_snapshot
    from app.schemas.selection import SelectionInput
    project = SimpleNamespace(id=uuid.uuid4(), name="Проект", parameters={"area": 999})
    source = snapshot()
    candidate = {**source["candidate"], "product_id": 7}
    original = {"parameters": {"area": 100}, "economics": {"scenarios": [{"name": "RaaS", "quantity": 2, "equipment": source["equipment"]}]},
        "bindings": [{"product_id": 7}], "selection": {"candidates": [candidate], "context": source["context"], "options": SelectionInput(process_id=1).model_dump()}}
    db = SimpleNamespace(scalar=AsyncMock(return_value=SimpleNamespace(inputs_snapshot=original)))
    data = SimulationRequest(project_updated_at=datetime.now(UTC), economics_run_id=uuid.uuid4())
    result = await simulation_snapshot(db, project, data)
    assert result["parameters"] == {"area": 100}
    assert result["quantity"] == 2
    bound = db.scalar.call_args.args[0].compile().params
    assert project.id in bound.values() and data.economics_run_id in bound.values()
    db.scalar.return_value = None
    with pytest.raises(HTTPException) as err: await simulation_snapshot(db, project, data)
    assert err.value.status_code == 404


@pytest.mark.asyncio
async def test_guest_cannot_save_and_stale_project_cannot_run(monkeypatch):
    from app.api.routes import simulation as route
    from app.schemas.selection import SelectionInput
    data = SimulationRequest(project_updated_at=datetime.now(UTC), selection=SelectionInput(process_id=1), product_id=1, save=True)
    with pytest.raises(HTTPException) as err: await route.calculate_simulation(uuid.uuid4(), data, None, None)
    assert err.value.status_code == 401
    data.save = False
    visible = AsyncMock(return_value=SimpleNamespace(updated_at=datetime(2020,1,1,tzinfo=UTC)))
    monkeypatch.setattr(route, "visible_project", visible)
    with pytest.raises(HTTPException) as err: await route.calculate_simulation(uuid.uuid4(), data, None, None)
    assert err.value.status_code == 409


@pytest.mark.asyncio
async def test_save_and_reopen_preserve_result_and_scope_access(monkeypatch):
    from app.api.routes import simulation as route
    from app.schemas.selection import SelectionInput
    from app.models.enums import CalculationType, CalculationStatus
    version = datetime.now(UTC)
    project = SimpleNamespace(id=uuid.uuid4(), updated_at=version, parameters={})
    user = SimpleNamespace(id=uuid.uuid4())
    source = {**snapshot(), "parameters": {}}
    visible = AsyncMock(return_value=project)
    monkeypatch.setattr(route, "visible_project", visible)
    monkeypatch.setattr(route, "simulation_snapshot", AsyncMock(return_value=source))
    records = []
    def add(record):
        record.id = uuid.uuid4()
        record.created_at = version
        records.append(record)
    db = SimpleNamespace(add=add, commit=AsyncMock(), scalar=AsyncMock())
    data = SimulationRequest(project_updated_at=version, selection=SelectionInput(process_id=1), product_id=7, save=True)
    response = await route.calculate_simulation(project.id, data, db, user)
    record = records[0]
    assert record.calc_type == CalculationType.SIMULATION
    assert record.status == CalculationStatus.SUCCEEDED
    assert record.created_by_id == user.id
    assert record.inputs_snapshot["options"] == data.options.model_dump()
    assert response["results"] == record.results
    assert visible.call_args.kwargs == {"write": True}
    db.commit.assert_awaited_once()
    # Reopening must return persisted numbers without executing the engine.
    monkeypatch.setattr(route, "simulate", lambda *args: pytest.fail("reopening must not simulate"))
    db.scalar.return_value = record
    reopened = await route.get_simulation(project.id, record.id, db, user)
    assert reopened["results"] == response["results"]
    params = db.scalar.call_args.args[0].compile().params
    assert project.id in params.values() and record.id in params.values()
    project.parameters = {"area": 10}
    assert (await route.get_simulation(project.id, record.id, db, user))["stale"]
    db.scalar.return_value = None
    with pytest.raises(HTTPException) as err: await route.get_simulation(project.id, uuid.uuid4(), db, user)
    assert err.value.status_code == 404


def test_shift_clipping_preserves_actual_motion_duration():
    source = snapshot()
    source["context"]["hours_per_day"] = .01
    source["context"]["daily_demand"] = 1
    result = simulate(source, SimulationOptions(route_m=100))
    last = result["robots"][0]["segments"][-1]
    assert last["state"] == "outbound"
    assert last["end"] == 36
    assert last["phase_end"] == 100


def test_large_queue_preserves_demand_without_old_12000_limit():
    source = snapshot()
    source["context"]["daily_demand"] = 12001
    result = simulate(source, SimulationOptions())
    assert result["kpi"]["target"] == 12001
    assert result["frames"][-1]["arrived"] == 12001
    assert result["kpi"]["completed"] <= 12001
