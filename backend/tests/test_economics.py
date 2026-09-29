import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.api.routes.economics import router
from app.schemas.economics import EconomicsInput
from app.services.economics import calculate


def inputs(**scenario):
    return EconomicsInput.model_validate({
        "baseline_annual_labor": 10000, "baseline_annual_other": 2000, "electricity_price": .1,
        "scenarios": [{"name": "Робот", "quantity": 2, "equipment_price": 1000, "software": 200,
                       "infrastructure": 100, "integration": 100, "training": 100, "reserve_percent": 10,
                       "labor_saving_percent": 50, "other_saving_percent": 25, "annual_service_per_robot": 100,
                       "annual_licenses": 100, "annual_other": 100, "annual_operators": 500,
                       "power_kw": .5, "annual_additional_benefit": 100, "service_life_years": 2, **scenario}],
    })


def test_purchase_costs_cashflow_and_replacements():
    response = calculate(inputs())
    r = response.results[0]
    assert response.baseline_tco == 60000
    assert (r.capex, r.annual_opex, r.annual_effect, r.tco) == (2750, 7600, 4500, 44750)
    assert [y.replacement for y in r.years] == [0, 0, 0, 2000, 0, 2000]
    assert r.net_effect == 15750 and r.roi_percent == 572.73
    assert r.simple_payback_years == .6111
    assert r.net_effect == sum(y.cashflow for y in r.years) == r.years[-1].cumulative
    assert sum(r.capex_breakdown.values()) == r.capex
    assert sum(r.opex_breakdown.values()) == r.annual_opex


def test_raas_has_rent_no_equipment_purchase_or_replacement():
    r = calculate(inputs(mode="raas", monthly_fee=100, annual_service_per_robot=0)).results[0]
    assert (r.capex, r.annual_opex, r.annual_effect) == (550, 9800, 2300)
    assert r.opex_breakdown["RaaS"] == 2400
    assert all(y.replacement == 0 for y in r.years)


def test_lifetime_equal_to_horizon_needs_no_replacement():
    r = calculate(inputs(service_life_years=5)).results[0]
    assert all(y.replacement == 0 for y in r.years)


def test_zero_investment_and_negative_effect_are_explicit():
    r = calculate(inputs(equipment_price=0, software=0, infrastructure=0, integration=0, training=0)).results[0]
    assert r.roi_percent is None and r.simple_payback_years is None
    r = calculate(inputs(annual_other=100000)).results[0]
    assert r.annual_effect < 0 and r.roi_percent < 0 and r.simple_payback_years is None


@pytest.mark.parametrize("patch", [{"equipment_price": None}, {"mode": "raas"}, {"quantity": 0},
                                  {"quantity": 1.5}, {"labor_saving_percent": 101}, {"power_kw": float("nan")},
                                  {"equipment_price": float("inf")}, {"service_life_years": 0}])
def test_invalid_assumptions_are_rejected(patch):
    with pytest.raises(ValidationError):
        inputs(**patch)


def test_guest_api_returns_snapshot_and_rejects_bad_inputs():
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    request = inputs().model_dump()
    response = client.post("/economics/calculate", json=request)
    assert response.status_code == 200
    assert response.json()["inputs"] == request
    assert response.json()["formulas"] and response.json()["assumptions"]
    request["horizon_years"] = 0
    assert client.post("/economics/calculate", json=request).status_code == 422


@pytest.mark.parametrize("mode", ["purchase", "raas"])
def test_unit_metrics_use_full_cost_and_same_volume(mode):
    data = inputs(mode=mode, monthly_fee=100)
    data.daily_volume = 10
    data.days_per_year = 200
    data.volume_unit = "паллета"
    response = calculate(data)
    r = response.results[0]
    assert response.annual_volume == 2000
    assert response.baseline_cost_per_unit == response.baseline_tco / 10000
    assert r.cost_per_unit == round(r.tco / 10000, 4)
    assert r.saving_per_unit == round((response.baseline_tco - r.tco) / 10000, 4)
    assert r.tco_saving_percent == round(100 * (response.baseline_tco - r.tco) / response.baseline_tco, 2)


def test_missing_volume_and_zero_baseline_are_not_fake_zero_metrics():
    data = inputs()
    data.baseline_annual_labor = data.baseline_annual_other = 0
    result = calculate(data)
    assert result.annual_volume is None and result.baseline_cost_per_unit is None
    assert result.results[0].cost_per_unit is None
    assert result.results[0].saving_per_unit is None
    assert result.results[0].tco_saving_percent is None


@pytest.mark.parametrize("volume", [0, -1, float("nan"), float("inf"), 1e-12])
def test_invalid_volume_is_rejected(volume):
    with pytest.raises(ValidationError):
        EconomicsInput.model_validate({**inputs().model_dump(), "daily_volume": volume})


def test_defaults_follow_reference_units_and_are_independent():
    from app.services.economics_defaults import default_profile
    profile = default_profile()
    assert profile["common"]["electricity_price"] == 9
    assert profile["common"]["baseline_annual_labor"] == 5 * 80000 * 12 * 1.302
    assert profile["scenario"]["reserve_percent"] == 10
    assert profile["scenario"]["component_replacement_interval"] == 4
    for mode in ["purchase", "raas"]:
        data = EconomicsInput(**profile["common"], scenarios=[{"name": "Пример", "mode": mode, **profile["scenario"]}])
        assert calculate(data).results[0].cost_per_unit is not None
    changed = default_profile([{"code":"capex_contingency","value":.18,"source":"Объект"}, {"code":"electricity_price","value":0,"source":"Договор"}])
    assert changed["scenario"]["reserve_percent"] == 18
    assert changed["common"]["electricity_price"] == 0
    assert changed["field_sources"]["common.electricity_price"] == "Договор"
    assert default_profile()["scenario"]["reserve_percent"] == 10
    invalid = default_profile([{"code":"battery_replacement_years","value":0,"source":"Некорректный"}])
    assert invalid["scenario"]["component_replacement_interval"] == 4
    assert any("вне допустимого" in n for n in invalid["notes"])


def test_sensitivity_scales_unit_volume_with_fleet():
    from app.services.economics_reports import sensitivity
    data = inputs()
    data.daily_volume = 100
    result = sensitivity(data, 20)
    plus = next(r for r in result["rows"] if r["factor"] == "volume" and r["delta_percent"] == 20)["results"][1]
    assert plus["cost_per_unit"] == round(plus["tco"] / (120 * data.days_per_year * data.horizon_years), 4)


def test_reports_contain_unit_metrics_and_sources():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.economics_reports import workbook_report, sensitivity
    from app.services.economics_pdf import pdf_report
    data = inputs()
    data.daily_volume = 100
    result = calculate(data).model_dump(mode="json")
    analysis = sensitivity(data, 20)
    book = load_workbook(BytesIO(workbook_report(result, analysis)))
    assert "Удельные затраты" in book.sheetnames
    assert book["Удельные затраты"].cell(3,2).value == result["results"][0]["cost_per_unit"]
    pdf = pdf_report(result, analysis)
    assert pdf.startswith(b"%PDF-") and len(pdf) > 10000
    # Historical snapshots without the new fields remain exportable.
    del result["baseline_cost_per_unit"]
    assert pdf_report(result).startswith(b"%PDF-")
