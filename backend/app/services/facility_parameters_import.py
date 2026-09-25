"""Импорт справочника параметров объектов из Excel организатора (Датасеты_хакатон.xlsx).

Лист = тип объекта. Строка «▌ …» открывает раздел, остальные строки — параметры:
название, единица, базовое значение, min, max, примечание (п. 3.2.1, 3.2.5 ТЗ).
Названиям параметров сопоставлены стабильные коды, на которые опираются расчётные модули.
"""

import hashlib
import io
import re
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import DatasetVersion, DataSource, FacilityType, ParameterDefinition
from app.models.enums import DatasetKind, SourceType, ValueDataType

SHEET_TO_FACILITY = {"Склад": "warehouse", "Аэропорт": "airport", "Медучреждение": "medical"}

# (код, обязательный) по названию параметра. Обязательность — по минимальному составу п. 3.2.1 ТЗ.
CODES: dict[str, dict[str, tuple[str, bool]]] = {
    "Склад": {
        "Общая площадь склада": ("total_area_m2", True),
        "Площадь активной (роботизируемой) зоны": ("active_area_m2", True),
        "Высота потолков в зоне хранения": ("ceiling_height_m", False),
        "Количество этажей (мезонинов)": ("floors_count", False),
        "Ширина главных проездов": ("main_aisle_width_m", True),
        "Ширина рабочих проходов между стеллажами": ("rack_aisle_width_m", True),
        "Тип напольного покрытия": ("floor_type", False),
        "Ровность пола (отклонение)": ("floor_flatness_mm_per_2m", False),
        "Количество рабочих смен в сутки": ("shifts_per_day", True),
        "Рабочих дней в году": ("working_days_per_year", True),
        "Продолжительность смены": ("shift_duration_h", True),
        "Пиковый коэффициент нагрузки": ("peak_load_factor", True),
        "Объём приёмки (поддоны/сутки)": ("inbound_pallets_per_day", True),
        "Объём отгрузки (поддоны/сутки)": ("outbound_pallets_per_day", True),
        "Объём отбора (строк/сутки, всего)": ("picking_lines_per_day", True),
        "Объём отбора (штук/сутки, всего)": ("picking_units_per_day", False),
        "Доля мелкоштучного отбора (piece-pick)": ("piece_pick_share_pct", False),
        "Количество SKU (активных)": ("sku_count", True),
        "Доля SKU с быстрым оборотом (A-класс)": ("sku_a_class_share_pct", False),
        "Общая численность персонала склада": ("staff_total", True),
        "Из них: отборщики (комплектовщики)": ("pickers_count", True),
        "Из них: операторы погрузчиков": ("forklift_operators_count", False),
        "Из них: операторы упаковочных линий": ("packing_operators_count", False),
        "Средняя з/п отборщика (gross)": ("picker_salary_rub_month", True),
        "Средняя з/п оператора погрузчика (gross)": ("forklift_operator_salary_rub_month", False),
        "Коэффициент начислений на ФОТ (страховые взносы)": ("payroll_tax_coef", True),
        "Средняя выработка отборщика (строк/ч)": ("picker_productivity_lines_h", True),
        "Коэффициент потерь рабочего времени (отпуск, болезнь, текучесть)": ("staff_time_loss_pct", False),
        "Средняя длина маршрута отборщика на 1 строку": ("picking_route_length_m", True),
        "Протяжённость конвейерной/транспортной системы": ("conveyor_length_m", False),
        "Тип стеллажной системы": ("storage_type", True),
        "Количество паллетомест": ("pallet_positions", False),
        "Средняя масса грузовой единицы (паллет)": ("pallet_weight_kg", True),
        "Средняя масса штучной единицы (SKU)": ("unit_weight_kg", False),
        "Средние габариты паллеты (Д×Ш×В)": ("pallet_dimensions_mm", True),
        "Средние габариты штучной единицы (Д×Ш×В)": ("unit_dimensions_mm", False),
        "Доля негабаритных/нестандартных грузов": ("oversize_share_pct", False),
        "Мощность электроснабжения (доступная)": ("available_power_kw", False),
        "Наличие WMS": ("has_wms", False),
        "Наличие ERP/1С": ("erp_system", False),
        "Планируемый бюджет на роботизацию (CAPEX)": ("capex_budget_mln_rub", False),
        "Горизонт расчёта окупаемости": ("horizon_years", True),
    },
    "Аэропорт": {
        "Суммарная площадь терминала (ов)": ("terminal_area_m2", True),
        "Площадь перрона и технических зон": ("apron_area_m2", True),
        "Количество терминалов": ("terminals_count", False),
        "Количество выходов на посадку (гейтов)": ("gates_count", False),
        "Количество взлётно-посадочных полос": ("runways_count", False),
        "Пассажиропоток (млн пассажиров/год)": ("annual_passengers_mln", True),
        "Среднесуточное количество пассажиров": ("daily_passengers", False),
        "Пиковое количество пассажиров в час (PHF)": ("peak_hour_passengers", True),
        "Доля трансферных пассажиров": ("transfer_share_pct", False),
        "Количество стоек регистрации": ("checkin_counters_count", False),
        "Среднесуточное количество рейсов (взлёт+посадка)": ("daily_flights", True),
        "Пиковое количество рейсов в час": ("peak_hour_flights", True),
        "Среднее время оборота воздушного судна (TAT)": ("turnaround_time_min", False),
        "Среднее количество операций наземного обслуживания на 1 рейс": ("ground_ops_per_flight", True),
        "Объём перемещения багажа (единиц/сутки)": ("baggage_units_per_day", True),
        "Средняя масса единицы багажа": ("baggage_unit_weight_kg", True),
        "Количество стоек выдачи багажа (каруселей)": ("baggage_carousels_count", False),
        "Объём бортового питания (порций/сутки)": ("catering_meals_per_day", False),
        "Объём заправки воздушных судов (рейсов/сут)": ("refueling_flights_per_day", False),
        "Суточное количество рейсов внутренних грузовых тележек (внутри терминала)": ("internal_cart_trips_per_day", False),
        "Количество уборочных машин (терминал)": ("cleaning_machines_count", False),
        "Площадь, убираемая роботизированной уборкой": ("robotic_cleaning_area_m2", False),
        "Суточный объём вывоза мусора (контейнеров)": ("waste_containers_per_day", False),
        "Численность персонала наземного обслуживания (рамп)": ("ramp_staff_count", True),
        "Численность персонала внутри терминала (логистика, уборка)": ("terminal_staff_count", True),
        "Средняя з/п сотрудника наземного обслуживания (gross)": ("ramp_staff_salary_rub_month", True),
        "Средняя з/п уборщика терминала (gross)": ("cleaner_salary_rub_month", False),
        "Коэффициент начислений на ФОТ": ("payroll_tax_coef", True),
        "Годовая текучесть (персонал терминала)": ("staff_turnover_pct", False),
        "Зонирование (количество режимных зон)": ("security_zones_count", True),
        "Наличие системы контроля доступа (СКУД)": ("access_control_system", False),
        "Требования по сертификации оборудования для airside": ("airside_certification", True),
        "Ограничения по уровню шума (зона)": ("noise_limit_dba", False),
        "Температура в неотапливаемых зонах (перрон, зима)": ("outdoor_min_temp_c", False),
        "Наличие FIDS/AODB системы": ("has_fids_aodb", False),
        "Наличие BMS (системы управления зданием)": ("has_bms", False),
        "Доступная мощность для зарядной инфраструктуры": ("charging_power_kw", False),
        "Планируемый бюджет на роботизацию (CAPEX)": ("capex_budget_mln_rub", False),
        "Горизонт расчёта окупаемости": ("horizon_years", True),
    },
    "Медучреждение": {
        "Тип медицинского учреждения": ("facility_subtype", True),
        "Общая площадь здания(й)": ("total_area_m2", True),
        "Количество этажей (основной корпус)": ("floors_count", True),
        "Количество лифтов (грузовых/медицинских)": ("elevators_count", True),
        "Количество коек (стационар)": ("beds_count", True),
        "Коечный фонд в эксплуатации (средняя занятость)": ("bed_occupancy_pct", False),
        "Количество операционных": ("operating_rooms_count", False),
        "Количество амбулаторных посещений в сутки": ("outpatient_visits_per_day", False),
        "Режим работы стационара": ("inpatient_schedule", True),
        "Режим работы амбулатории": ("outpatient_schedule", False),
        "Количество смен медперсонала (уход за пациентами)": ("medical_shifts_per_day", False),
        "Пиковое время логистической нагрузки": ("logistics_peak_hours", False),
        "Количество кормлений в сутки": ("meals_per_day_per_patient", False),
        "Общее количество порций питания в сутки": ("meal_portions_per_day", True),
        "Среднее расстояние от пищеблока до отделения": ("kitchen_to_ward_distance_m", True),
        "Количество точек раздачи питания (отделений)": ("food_delivery_points", False),
        "Средняя масса тележки с питанием (брутто)": ("food_cart_weight_kg", False),
        "Норматив доставки питания (мин от пищеблока до отделения)": ("food_delivery_norm_min", False),
        "Объём грязного белья (кг/сутки)": ("dirty_linen_kg_per_day", True),
        "Объём чистого белья на раздачу (кг/сутки)": ("clean_linen_kg_per_day", False),
        "Количество точек сбора/выдачи белья": ("linen_points", False),
        "Периодичность смены белья (раз в сутки, в среднем)": ("linen_change_per_day", False),
        "Средняя масса контейнера с бельём": ("linen_container_weight_kg", False),
        "Количество наименований медикаментов в обращении": ("medicine_items_count", False),
        "Объём выдачи медикаментов (заявок/сутки)": ("medicine_requests_per_day", True),
        "Количество аптечных точек выдачи (аптека, аптечные склады)": ("pharmacy_points", False),
        "Количество точек доставки (отделений + ОР + реанимация)": ("medicine_delivery_points", False),
        "Среднее время комплектации 1 заявки в аптеке": ("pharmacy_picking_time_min", False),
        "Доля срочных (STAT) доставок медикаментов": ("stat_delivery_share_pct", False),
        "Объём доставки расходных материалов (рейсов/сутки)": ("consumables_trips_per_day", True),
        "Количество биоматериалов (проб) в сутки": ("lab_samples_per_day", True),
        "Количество клинико-диагностических лабораторий (КДЛ)": ("labs_count", False),
        "Среднее время доставки пробы (норматив)": ("sample_delivery_norm_min", False),
        "Объём выдачи результатов анализов (рейсов/сутки)": ("lab_results_trips_per_day", False),
        "Объём медицинских отходов класса А (ненасыщенные)": ("waste_class_a_kg_per_day", True),
        "Объём медицинских отходов класса Б (инфицированные)": ("waste_class_b_kg_per_day", True),
        "Количество точек сбора отходов": ("waste_points", False),
        "Периодичность вывоза отходов из отделений": ("waste_removal_per_day", False),
        "Численность санитаров и транспортировщиков": ("orderlies_count", True),
        "Численность сотрудников пищеблока (раздача)": ("kitchen_staff_count", False),
        "Численность сотрудников прачечной (транспорт белья)": ("laundry_staff_count", False),
        "Средняя з/п санитара/транспортировщика (gross)": ("orderly_salary_rub_month", True),
        "Средняя з/п сотрудника пищеблока (gross)": ("kitchen_staff_salary_rub_month", False),
        "Коэффициент начислений на ФОТ": ("payroll_tax_coef", True),
        "Годовая текучесть (немедицинский персонал)": ("staff_turnover_pct", False),
        "Обеззараживание робота между рейсами": ("disinfection_requirement", True),
        "Требования к уровню шума в палатах (ночное время)": ("ward_noise_limit_dba", False),
        "Наличие СКУД (контроль доступа по зонам)": ("access_control_system", True),
        "Требования к материалу поверхностей робота": ("robot_surface_material", False),
        "Наличие МИС (медицинская информационная система)": ("mis_system", False),
        "Наличие ЛИС (лабораторная информационная система)": ("has_lis", False),
        "Наличие системы управления лифтами (BMS)": ("elevator_bms", False),
        "Ширина коридоров (основных)": ("corridor_width_m", True),
        "Наличие пандусов/подъёмников (для межэтажного AMR без лифта)": ("has_ramps_lifts", False),
        "Доступная мощность для зарядной инфраструктуры": ("charging_power_kw", False),
        "Планируемый бюджет на роботизацию (CAPEX)": ("capex_budget_mln_rub", False),
        "Горизонт расчёта окупаемости": ("horizon_years", True),
    },
}

TEAM_SOURCE_NOTE = "Допущение команды: параметр требуется п. 3.2.1 ТЗ, но отсутствует в датасете организатора"

# Параметры из п. 3.2.1 ТЗ, которых нет в датасете организатора.
TEAM_ADDED: dict[str, list[dict[str, Any]]] = {
    "airport": [
        {"section": "Режим работы", "code": "operating_hours_per_day", "name": "Продолжительность работы в сутки",
         "unit": "ч", "data_type": "number", "default_value": 24, "min_value": 8, "max_value": 24,
         "is_required": True, "hint": "Режим работы объекта (п. 3.2.1 ТЗ). Крупные аэропорты работают круглосуточно."},
        {"section": "Режим работы", "code": "working_days_per_year", "name": "Рабочих дней в году", "unit": "дн.",
         "data_type": "integer", "default_value": 365, "min_value": 250, "max_value": 366, "is_required": True,
         "hint": "Аэропорт работает без выходных."},
        {"section": "Маршруты", "code": "avg_route_length_m",
         "name": "Средняя протяжённость маршрута робота (в одну сторону)", "unit": "м", "data_type": "number",
         "default_value": 600, "min_value": 50, "max_value": 5000, "is_required": True,
         "hint": "Протяжённость маршрутов (п. 3.2.1 ТЗ). Уточняется по схеме терминала."},
    ],
}

ENUMS = {
    ("warehouse", "storage_type"): ["Фронтальные паллетные", "Shuttle", "AutoStore", "Miniload", "Drive-in", "Push-back"],
    ("medical", "facility_subtype"): ["Поликлиника", "Многопрофильная больница", "Онкоцентр", "Диагностический центр"],
}

COUNT_UNITS = {"шт.", "чел.", "SKU", "смен", "дн.", "коек", "наименований", "м/п", "смен/сут"}
DIMENSIONS_RE = re.compile(r"\d+\s*[×x]\s*\d+\s*[×x]\s*\d+")


class FacilityWorkbookError(ValueError):
    """Файл не соответствует шаблону датасета объектов."""


@dataclass
class ParameterImportStats:
    status: str = "imported"
    dataset_version_id: int | None = None
    created: int = 0
    updated: int = 0
    by_facility: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _number(value: Any) -> float | int | None:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _detect_type(value: Any, unit: str | None) -> tuple[str, Any]:
    if isinstance(value, bool):
        return "boolean", value
    if _number(value) is not None:
        return ("integer" if (unit or "").strip() in COUNT_UNITS else "number"), value
    text = str(value).strip()
    if text in ("Да", "Нет"):
        return "boolean", text == "Да"
    if DIMENSIONS_RE.fullmatch(text):
        return "dimensions", text
    return "string", text


def _section_name(raw: str) -> str:
    text = raw.replace("▌", "").strip()
    return text[:1] + text[1:].lower()


def parse_facility_workbook(content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Разбирает Excel в список определений параметров. Чистая функция — не обращается к БД."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl бросает разные исключения на битых файлах
        raise FacilityWorkbookError("Не удалось открыть Excel-файл. Проверьте, что это .xlsx из набора организатора.") from exc

    missing_sheets = [s for s in SHEET_TO_FACILITY if s not in workbook.sheetnames]
    if missing_sheets:
        raise FacilityWorkbookError(f"В файле нет листов: {', '.join(missing_sheets)}. Ожидается шаблон «Датасеты_хакатон.xlsx».")

    result: list[dict[str, Any]] = []
    warnings: list[str] = []
    for sheet, facility in SHEET_TO_FACILITY.items():
        codes = CODES[sheet]
        section, order, seen = None, 0, set()
        for row in workbook[sheet].iter_rows(min_row=3, values_only=True):
            name, unit, value, vmin, vmax, note = (list(row) + [None] * 6)[:6]
            if name is None:
                continue
            name = str(name).strip()
            if name.startswith("▌"):
                section = _section_name(name)
                continue
            if name not in codes:
                warnings.append(f"Лист «{sheet}»: неизвестный параметр «{name}» пропущен — добавьте его код в CODES.")
                continue
            code, required = codes[name]
            seen.add(name)
            unit_text = None if unit in (None, "-") else str(unit).strip()
            data_type, default = _detect_type(value, unit_text)
            allowed = ENUMS.get((facility, code))
            order += 10
            result.append({
                "facility_type": facility, "code": code, "name": name, "section": section, "unit": unit_text,
                "data_type": "enum" if allowed else data_type, "is_required": required, "default_value": default,
                "min_value": _number(vmin), "max_value": _number(vmax), "allowed_values": allowed,
                "hint": str(note).strip() if note else None,
                "example": str(value).strip() if value is not None else None,
                "source": "organizer_facility_datasets", "source_note": None, "sort_order": order,
            })
        for name in sorted(set(codes) - seen):
            warnings.append(f"Лист «{sheet}»: нет параметра «{name}».")
        for extra in TEAM_ADDED.get(facility, []):
            order += 10
            result.append({
                "facility_type": facility, "allowed_values": None, "example": str(extra["default_value"]),
                "source": "team_assumption", "source_note": TEAM_SOURCE_NOTE, "sort_order": order, **extra,
            })
    workbook.close()
    return result, warnings


async def import_facility_parameters(session: AsyncSession, content: bytes, file_name: str) -> ParameterImportStats:
    """Загружает определения параметров (upsert по типу объекта и коду). Коммит — на стороне вызывающего кода."""
    stats = ParameterImportStats()
    checksum = hashlib.sha256(content).hexdigest()
    existing_version = await session.scalar(
        select(DatasetVersion).where(
            DatasetVersion.kind == DatasetKind.FACILITY_PARAMETERS, DatasetVersion.checksum_sha256 == checksum
        )
    )
    if existing_version:
        stats.status = "skipped"
        stats.dataset_version_id = existing_version.id
        return stats

    items, stats.warnings = parse_facility_workbook(content)
    facilities = {f.code: f for f in (await session.scalars(select(FacilityType))).all()}
    missing = sorted({i["facility_type"] for i in items} - set(facilities))
    if missing:
        raise FacilityWorkbookError(
            f"В БД нет типов объектов: {', '.join(missing)}. Сначала запустите backend — он создаёт базовые справочники."
        )
    sources = {
        s.code: s
        for s in (await session.scalars(select(DataSource).where(DataSource.code.in_(["organizer_facility_datasets", "team_assumption"])))).all()
    }
    if "organizer_facility_datasets" not in sources:
        sources["organizer_facility_datasets"] = DataSource(
            code="organizer_facility_datasets", title=f"Демо-датасеты объектов ({file_name})",
            source_type=SourceType.ORGANIZER, publisher="ФЦ БАС",
        )
        session.add(sources["organizer_facility_datasets"])
    if "team_assumption" not in sources:
        sources["team_assumption"] = DataSource(
            code="team_assumption", title="Допущение команды", source_type=SourceType.TEAM_ASSUMPTION
        )
        session.add(sources["team_assumption"])
    await session.flush()

    definitions = {
        (d.facility_type_id, d.code): d for d in (await session.scalars(select(ParameterDefinition))).all()
    }
    for item in items:
        facility = facilities[item["facility_type"]]
        values = {
            "name": item["name"], "section": item["section"], "unit": item["unit"],
            "data_type": ValueDataType(item["data_type"]), "is_required": item["is_required"],
            "default_value": item["default_value"], "min_value": item["min_value"], "max_value": item["max_value"],
            "allowed_values": item["allowed_values"], "hint": item["hint"], "example": item["example"],
            "source_id": sources[item["source"]].id, "source_note": item["source_note"], "sort_order": item["sort_order"],
        }
        definition = definitions.get((facility.id, item["code"]))
        if definition is None:
            session.add(ParameterDefinition(facility_type_id=facility.id, code=item["code"], **values))
            stats.created += 1
        else:
            for key, value in values.items():
                setattr(definition, key, value)
            stats.updated += 1
        stats.by_facility[facility.code] = stats.by_facility.get(facility.code, 0) + 1

    await session.execute(
        update(DatasetVersion).where(DatasetVersion.kind == DatasetKind.FACILITY_PARAMETERS).values(is_current=False)
    )
    version = DatasetVersion(
        kind=DatasetKind.FACILITY_PARAMETERS, label=file_name.rsplit(".", 1)[0], file_name=file_name,
        checksum_sha256=checksum, source_id=sources["organizer_facility_datasets"].id, row_count=len(items),
        stats={"created": stats.created, "updated": stats.updated, "by_facility": stats.by_facility}, is_current=True,
    )
    session.add(version)
    await session.flush()
    stats.dataset_version_id = version.id
    return stats
