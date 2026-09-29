import math
import csv
import io
import zipfile
import re
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from typing import Any, Sequence

from app.api.errors import ApiValidationError, FieldError
from app.models import ParameterDefinition


def parse_parameter_file(content: bytes, suffix: str, definitions: Sequence[ParameterDefinition]) -> dict[str, Any]:
    """Шаблон code,value,...,unit. Пустая ячейка очищает значение; отсутствующая строка оставляет его."""
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text), delimiter=";" if ";" in text.splitlines()[0] else ","))
        elif suffix == ".xlsx":
            import openpyxl
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20_000_000:
                    raise ValueError("Слишком большой распакованный Excel-файл.")
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet.max_row and sheet.max_row > 301 or sheet.max_column and sheet.max_column > 10:
                    raise ValueError("В шаблоне допускается до 300 параметров и 10 колонок.")
                rows = []
                for row in sheet.iter_rows(max_row=302, max_col=10):
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("Замените формулы Excel их значениями.")
                    rows.append([cell.value for cell in row])
                while rows and all(value is None for value in rows[-1]):
                    rows.pop()
            finally:
                workbook.close()
        else:
            raise ValueError("Поддерживаются CSV UTF-8 и XLSX.")
        if not rows or [str(v or "").strip().lower() for v in rows[0][:2]] != ["code", "value"]:
            raise ValueError("Первые колонки должны называться code и value. Скачайте шаблон проекта.")
        if len(rows) > 301:
            raise ValueError("В файле допускается до 300 параметров.")
        by_code = {d.code: d for d in definitions}
        headers = [str(v or "").strip().lower() for v in rows[0]]
        if headers.count("unit") != 1:
            raise ValueError("В шаблоне нужна колонка unit с единицами измерения. Скачайте актуальный шаблон.")
        unit_index = headers.index("unit")
        result = {}
        for row in rows[1:]:
            if not row or all(v in (None, "") for v in row):
                continue
            code = str(row[0] or "").strip()
            if len(row) < 2 or code not in by_code or code in result:
                raise ValueError(f"Неизвестный, повторный или незаполненный код параметра: {code}.")
            value = row[1]
            expected_unit = by_code[code].unit or ""
            actual_unit = str(row[unit_index] or "").strip() if len(row) > unit_index else ""
            normalize_unit = lambda unit: unit.replace("²", "2").replace("³", "3").replace(" ", "").lower().rstrip(".")
            if normalize_unit(actual_unit) != normalize_unit(expected_unit):
                raise ValueError(f"Параметр {code}: единица должна быть «{expected_unit or 'без единицы'}», получено «{actual_unit or 'пусто'}». Пересчитайте значение в единицы шаблона.")
            if isinstance(value, str):
                value = value.strip()
                if value.startswith("'") and value[1:2] in ("=", "+", "-", "@"):
                    value = value[1:]
            if value not in (None, ""):
                d = by_code[code]
                if d.data_type in ("number", "integer") and isinstance(value, str):
                    try:
                        value = float(value.replace(" ", "").replace("\u00a0", "").replace(",", "."))
                    except ValueError:
                        raise ValueError(f"Параметр {code}: введите число.")
                elif d.data_type == "boolean" and isinstance(value, str):
                    flags = {"true": True, "false": False, "да": True, "нет": False, "1": True, "0": False}
                    if value.lower() not in flags:
                        raise ValueError(f"Параметр {code}: укажите да или нет.")
                    value = flags[value.lower()]
            result[code] = value
        validate_parameters(result, definitions)
        return result
    except (ValueError, IndexError, UnicodeError, csv.Error, zipfile.BadZipFile, KeyError, ParseError, InvalidFileException) as exc:
        raise ApiValidationError([FieldError("file", str(exc))]) from exc


def parameter_template(parameters: dict[str, Any], definitions: Sequence[ParameterDefinition]) -> str:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["code", "value", "name", "unit", "required", "min", "max", "allowed_values", "default", "source"])
    for d in definitions:
        value = parameters.get(d.code, "")
        if isinstance(value, bool):
            value = "да" if value else "нет"
        cells = [d.code, value, d.name, d.unit or "", "да" if d.is_required else "нет", d.min_value, d.max_value,
                 " | ".join(map(str, d.allowed_values or [])), d.default_value,
                 d.source.title if d.source else d.source_note or "Не указан"]
        writer.writerow(["'" + c if isinstance(c, str) and c.startswith(("=", "+", "-", "@")) else c for c in cells])
    return "\ufeff" + output.getvalue()


def parameter_workbook(parameters, definitions) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Параметры"
    rows = list(csv.reader(io.StringIO(parameter_template(parameters, definitions).lstrip("\ufeff")), delimiter=";"))
    for row in rows:
        sheet.append(row)
    for index, d in enumerate(definitions, start=2):
        cell = sheet.cell(index, 2)
        value = parameters.get(d.code)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.value = value
        if d.is_required:
            cell.fill = PatternFill("solid", fgColor="EAF2FF")
        validation = None
        if d.data_type == "boolean":
            validation = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
        elif d.data_type == "enum" and len(",".join(map(str, d.allowed_values or []))) < 250 and not any('"' in str(v) or ',' in str(v) for v in d.allowed_values or []):
            validation = DataValidation(type="list", formula1='"'+','.join(map(str, d.allowed_values or []))+'"', allow_blank=True)
        elif d.data_type in ("number", "integer"):
            validation = DataValidation(type="whole" if d.data_type == "integer" else "decimal", operator="between",
                                        formula1=float(d.min_value) if d.min_value is not None else -1e15,
                                        formula2=float(d.max_value) if d.max_value is not None else 1e15, allow_blank=True)
        if validation:
            validation.errorTitle = "Проверьте значение"
            validation.error = "Используйте тип и диапазон из шаблона."
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(cell)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2463D8")
    for col, width in {"A": 36, "B": 25, "C": 60, "D": 18, "E": 12, "F": 15, "G": 15, "H": 45, "I": 25, "J": 55}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    guide = workbook.create_sheet("Инструкция")
    for text in ["Заполняйте колонку value на листе Параметры.", "Не меняйте code и unit. Единицы проверяются при загрузке.",
                 "Поля required=да обязательны для расчёта. Черновик можно сохранить неполным.",
                 "Пустое value очищает поле; удалённая строка оставляет прежнее значение.",
                 "default — справочное значение, source — его источник. Они не подставляются автоматически.",
                 "Вставляйте значения, а не формулы. Размер файла — до 2 МБ."]:
        guide.append([text])
    guide.column_dimensions["A"].width = 110
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def validate_parameters(values: dict[str, Any], definitions: Sequence[ParameterDefinition], *, require_complete=False) -> dict[str, Any]:
    """Черновики допускают незаполненные поля; введённые значения строго проверяются."""
    by_code = {d.code: d for d in definitions}
    errors: list[FieldError] = []
    result: dict[str, Any] = {}
    for code, value in values.items():
        d = by_code.get(code)
        error = None
        if d is None:
            errors.append(FieldError(f"parameters.{code}", "Неизвестный параметр для этого типа объекта."))
            continue
        if value is None or value == "":
            continue
        if d.data_type in ("number", "integer"):
            if isinstance(value, bool) or not isinstance(value, (int, float)) or abs(value) > 1e15 or not math.isfinite(value):
                error = "Введите конечное число."
            elif d.data_type == "integer" and value != int(value):
                error = "Введите целое число."
            elif d.min_value is not None and value < float(d.min_value):
                error = f"Минимальное значение: {d.min_value}."
            elif d.max_value is not None and value > float(d.max_value):
                error = f"Максимальное значение: {d.max_value}."
        elif d.data_type == "boolean":
            if not isinstance(value, bool):
                error = "Выберите да или нет."
        elif d.data_type == "enum":
            if value not in (d.allowed_values or []):
                error = "Выберите значение из списка."
        elif not isinstance(value, str) or len(value) > 2000:
            error = "Введите текст не длиннее 2000 символов."
        elif isinstance(value, str):
            value = value.strip()
            if not value:
                continue
            if d.data_type == "dimensions" or code.endswith("dimensions_mm"):
                parts = re.split(r"\s*[xх×*]\s*", value.lower())
                try:
                    if len(parts) != 3 or any(not 0 < float(p.replace(',', '.')) <= 100000 for p in parts):
                        raise ValueError()
                except ValueError:
                    error = "Укажите Д×Ш×В в мм, например 1200×800×1500; каждое измерение больше 0 и не более 100000."
        if error:
            errors.append(FieldError(f"parameters.{code}", error))
        else:
            result[code] = value
    for smaller, larger in (("active_area_m2", "total_area_m2"), ("available_area_m2", "total_area_m2"), ("pickers_count", "staff_total")):
        if isinstance(result.get(smaller), (int, float)) and isinstance(result.get(larger), (int, float)) and result[smaller] > result[larger]:
            errors.append(FieldError(f"parameters.{smaller}", f"Значение не может превышать «{by_code[larger].name}»."))
    if "shifts_per_day" in result and "shift_duration_h" in result and result["shifts_per_day"] * result["shift_duration_h"] > 24:
        errors.append(FieldError("parameters.shift_duration_h", "Суммарная продолжительность смен не может превышать 24 часа в сутки."))
    if require_complete:
        errors.extend(FieldError(f"parameters.{code}", "Обязательное поле для расчёта.") for code in missing_required(result, definitions))
    if errors:
        raise ApiValidationError(errors)
    return result


def missing_required(values: dict[str, Any], definitions: Sequence[ParameterDefinition]) -> list[str]:
    return [d.code for d in definitions if d.is_required and (d.code not in values or values[d.code] in (None, ""))]
