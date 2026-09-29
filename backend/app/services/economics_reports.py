"""Отчёты и однофакторная чувствительность на основе общей экономической модели."""
from io import BytesIO
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.schemas.economics import EconomicsInput
from app.services.economics import calculate


FACTORS = {
    "labor": "Расходы базового процесса на персонал",
    "price": "Цена оборудования / ставка RaaS",
    "service": "Ежегодное обслуживание при покупке",
    "volume": "Объём операций и требуемый парк",
}


def sensitivity(data: EconomicsInput, spread: float):
    base = calculate(data)
    rows = []
    for factor, label in FACTORS.items():
        for delta in (-spread, 0, spread):
            variant = data.model_dump()
            multiplier = 1 + delta / 100
            if factor == "labor":
                variant["baseline_annual_labor"] *= multiplier
            elif factor != "volume":
                for scenario in variant["scenarios"]:
                    field = ("equipment_price" if scenario["mode"] == "purchase" else "monthly_fee") if factor == "price" else "annual_service_per_robot"
                    if factor == "service" and scenario["mode"] == "raas":
                        continue
                    if scenario[field] is not None:
                        scenario[field] *= multiplier
            # Производные варианты допускают увеличение денежных величин сверх лимита формы.
            # Исходные значения уже проверены схемой, multiplier ограничен диапазоном 0,5–1,5.
            adjusted = data.model_copy(deep=True)
            adjusted.baseline_annual_labor = variant["baseline_annual_labor"]
            for target, values in zip(adjusted.scenarios, variant["scenarios"]):
                target.equipment_price = values["equipment_price"]
                target.monthly_fee = values["monthly_fee"]
                target.annual_service_per_robot = values["annual_service_per_robot"]
            if factor == "volume":
                for target in adjusted.scenarios:
                    target.quantity = max(1, math.ceil((target.fleet_unrounded or target.quantity) * multiplier)) if delta else target.quantity
                    if target.equipment:
                        target.equipment.peak_rate *= multiplier
            result = base if delta == 0 else calculate(adjusted)
            rows.append({"factor": factor, "label": label, "delta_percent": delta,
                         "results": [{"name": "Без роботизации", "tco": result.baseline_tco,
                                      "net_effect": 0, "simple_payback_years": None, "roi_percent": None}] + result.model_dump(mode="json")["results"]})
    return {"model_version": base.model_version, "spread_percent": spread, "rows": rows,
            "assumptions": ["Каждый фактор меняется отдельно. Для объёма операций пересчитываются парк, зарядки и рабочие посты; для остальных факторов парк фиксирован.",
                "Это диапазон допущений, а не прогноз или доверительный интервал.",
                "Изменение нулевой статьи не влияет на результат. Обслуживание RaaS включено в ставку.",
                "При изменении объёма парк = ceil(исходное неокруглённое количество × множитель объёма). Если точного значения нет, используется указанное количество как допущение. Фиксированные ручные количества зарядок и постов сохраняются. ФОТ, доли экономии и постоянные расходы не масштабируются без отдельной модели базового процесса."]}


def workbook_report(result: dict, analysis: dict | None = None, snapshot: dict | None = None, title="Экономика роботизации") -> bytes:
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
        for row in ws:
            for cell in row:
                # Пользовательские названия и источники — текст, никогда формулы Excel.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="193D54")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(65, max(18, max(len(str(c.value or "")) for c in col[:100]) + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    sheet("Отчёт", ["Поле", "Значение"], [["Название", title], ["Версия модели", result["model_version"]],
        ["Горизонт, лет", result["inputs"]["horizon_years"]], ["Валюта", "RUB"],
        ["OPEX базы, ₽/год", result["baseline_annual_opex"]], ["TCO базы, ₽", result["baseline_tco"]],
        ["Примечание", "Альтернативные сценарии одного процесса; их эффект не суммируется. Пустая окупаемость/ROI означает, что показатель не определён."]])
    keys = ["name", "mode", "capex", "annual_opex", "annual_effect", "tco", "net_effect", "simple_payback_years", "roi_percent"]
    sheet("Сценарии", ["Сценарий", "Модель", "CAPEX, ₽", "OPEX, ₽/год", "Эффект, ₽/год", "TCO, ₽", "Чистый эффект, ₽", "Окупаемость, лет", "ROI, %"],
          [["Без роботизации", "baseline", 0, result["baseline_annual_opex"], 0, result["baseline_tco"], 0, None, None]] + [[r[k] for k in keys] for r in result["results"]])
    sheet("Интерпретация", ["Сценарий", "Экономия ФОТ, ₽/год", "Изменение OPEX, ₽/год", "Вывод", "Риски"],
          [[r["name"], r.get("annual_labor_saving"), r.get("annual_opex_change"), r.get("interpretation", ""), "\n".join(r.get("risks", []))] for r in result["results"]])
    sheet("Денежные потоки", ["Сценарий", "Год", "OPEX, ₽", "Замены, ₽", "Поток, ₽", "Накоплено, ₽"],
          [[r["name"], y["year"], y["opex"], y["replacement"], y["cashflow"], y["cumulative"]] for r in result["results"] for y in r["years"]])
    sheet("Статьи затрат", ["Сценарий", "Раздел", "Статья", "Сумма, ₽"],
          [[r["name"], section, key, value] for r in result["results"] for section in ("capex_breakdown", "opex_breakdown") for key, value in r[section].items()])
    equipment_rows = [[r["name"], i["name"], i["calculated"], i["quantity"], i["unit_price"], i["cost"]]
                      for r in result["results"] if r.get("equipment") for i in r["equipment"]["items"]]
    if equipment_rows:
        sheet("Состав оборудования", ["Сценарий", "Позиция", "Расчёт, шт.", "Принято, шт.", "Цена, ₽/шт.", "Стоимость, ₽"], equipment_rows)
        sheet("Допущения оборудования", ["Сценарий", "Описание"],
              [[r["name"], text] for r in result["results"] if r.get("equipment")
               for text in r["equipment"]["formulas"] + r["equipment"]["assumptions"]])
    rows = []
    def flatten(value, path=""):
        if isinstance(value, dict):
            for key, item in value.items(): flatten(item, f"{path}.{key}" if path else str(key))
        elif isinstance(value, list):
            for index, item in enumerate(value): flatten(item, f"{path}[{index}]")
        else:
            rows.append([path, value])
    flatten(snapshot or result["inputs"])
    sheet("Исходные данные", ["Параметр", "Значение"], rows)
    sheet("Формулы и допущения", ["Показатель", "Описание"], list(result["formulas"].items()) + [["Допущение", a] for a in result["assumptions"]])
    if analysis:
        sheet("Чувствительность", ["Фактор", "Изменение, %", "Сценарий", "TCO, ₽", "Чистый эффект, ₽", "Окупаемость, лет", "ROI, %"],
              [[row["label"], row["delta_percent"], r["name"], r["tco"], r["net_effect"], r["simple_payback_years"], r["roi_percent"]] for row in analysis["rows"] for r in row["results"]])
        sheet("Метод чувствительности", ["Примечание"], [[a] for a in analysis["assumptions"]])
    output = BytesIO()
    book.save(output)
    return output.getvalue()
